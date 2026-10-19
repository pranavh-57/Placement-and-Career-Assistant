from __future__ import annotations
import io
from app.utils.pdf_parser import extract_text_from_pdf

SUPPORTED_EXTENSIONS = {"pdf", "docx", "txt"}


def extract_text_from_docx(file_bytes: bytes) -> str:
    from docx import Document  # python-docx

    doc = Document(io.BytesIO(file_bytes))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]

    # Also pull text out of tables, since resumes sometimes use them for layout
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)

    return "\n".join(parts)


def extract_text_from_txt(file_bytes: bytes) -> str:
    return file_bytes.decode("utf-8", errors="ignore")


def extract_text(filename: str, file_bytes: bytes) -> str:
    """Dispatch to the right extractor based on file extension."""
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""

    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type: .{ext}. Supported types: "
            f"{', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    if ext == "pdf":
        return extract_text_from_pdf(file_bytes)
    elif ext == "docx":
        return extract_text_from_docx(file_bytes)
    elif ext == "txt":
        return extract_text_from_txt(file_bytes)